"""Rebuild SQLite with full feature extraction + curated table scaffolds.

Layers in `streets`:
  artera_raw       - original
  street_type      - 'Strada', 'Aleea', ...
  name             - display form (post-1993 diacritics, original â/î preserved)
  name_normalized  - lowercase ASCII, î≡â collapsed (matching key)
  title            - 'Prof.', 'Învățător', ...
  rank             - 'General', 'Voievod', ...
  is_saint, is_date, is_numeric  (0/1 flags)
  core_name        - name with title/rank/saint prefixes stripped
  core_name_norm   - normalized version of core_name (the join key for `persons`)
"""
import openpyxl, sqlite3, re
from pathlib import Path

from streets_lib import (
    DIACRITIC_FIX, fix_diacritics, normalize_match, STREET_TYPES,
    parse_artery, extract_features,
)

import argparse
_ap = argparse.ArgumentParser()
_ap.add_argument("--limit", type=int, default=None)
_ap.add_argument("--src", default="data/reference/registrul sectiilor de vot 14.05.2025.xlsx")
_ap.add_argument("--db",  default="data/streets.db")
_args = _ap.parse_args()

SRC = Path(_args.src)
DB  = Path(_args.db)
LIMIT = _args.limit
if DB.exists(): DB.unlink()

# normalization helpers (DIACRITIC_FIX, fix_diacritics, normalize_match),
# STREET_TYPES, and feature extraction (parse_artery, extract_features) live
# in streets_lib so the OSM/postal enrichment tools can reuse them.

# ---------- schema ----------
con = sqlite3.connect(DB)
con.executescript("""
CREATE TABLE streets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    judet TEXT NOT NULL,
    uat TEXT NOT NULL,
    siruta INTEGER,
    artera_raw TEXT NOT NULL,
    street_type TEXT,
    name TEXT,
    name_normalized TEXT,
    title TEXT,
    rank TEXT,
    is_saint INTEGER NOT NULL DEFAULT 0,
    is_date INTEGER NOT NULL DEFAULT 0,
    is_numeric INTEGER NOT NULL DEFAULT 0,
    core_name TEXT,
    core_name_norm TEXT
);
CREATE TABLE street_aliases (
    street_id INTEGER NOT NULL REFERENCES streets(id),
    alias TEXT NOT NULL,
    alias_normalized TEXT NOT NULL
);

-- ===== Curated lookup tables (scaffolds) =====
CREATE TABLE name_categories (
    core_name_norm TEXT PRIMARY KEY,
    category TEXT NOT NULL,
    subcategory TEXT,
    notes TEXT
);
CREATE TABLE persons (
    core_name_norm TEXT PRIMARY KEY,
    full_name TEXT NOT NULL,
    gender TEXT,
    birth_year INTEGER,
    death_year INTEGER,
    era TEXT,
    profession TEXT,
    nationality TEXT DEFAULT 'RO',
    wikidata_qid TEXT,
    wiki_sitelinks INTEGER,
    wiki_scope TEXT,
    wiki_ro_url TEXT,
    wiki_en_url TEXT,
    wiki_ro_views INTEGER,
    birth_place_qid TEXT,
    birth_place_label TEXT,
    birth_judet TEXT,
    cause_of_death_qid TEXT,
    cause_of_death_label TEXT,
    notes TEXT
);
CREATE TABLE place_refs (
    core_name_norm TEXT PRIMARY KEY,
    place_name TEXT NOT NULL,
    place_type TEXT,
    country TEXT DEFAULT 'RO',
    notes TEXT
);
CREATE TABLE nature_terms (
    core_name_norm TEXT PRIMARY KEY,
    term TEXT NOT NULL,
    nature_type TEXT,
    notes TEXT
);

-- ===== OSM enrichment scaffolds (populated by tools/osm_*.py) =====
-- One logical street per (uat_siruta, name_normalized). OSM splits a single
-- street into many `way` rows at every junction; tools/osm_ingest.py groups
-- those before insert, so this table mirrors streets_dedup's grain.
CREATE TABLE osm_streets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    uat_siruta INTEGER NOT NULL,
    name TEXT NOT NULL,
    name_normalized TEXT NOT NULL,
    street_type TEXT,               -- best-effort, via streets_lib.strip_street_type
    title TEXT,
    rank TEXT,
    is_saint INTEGER NOT NULL DEFAULT 0,
    is_date INTEGER NOT NULL DEFAULT 0,
    is_numeric INTEGER NOT NULL DEFAULT 0,
    core_name TEXT,
    core_name_norm TEXT,
    highway_class TEXT NOT NULL,
    ref TEXT,
    length_m REAL NOT NULL,
    way_ids TEXT NOT NULL,         -- JSON array of contributing OSM way ids
    geometry_wkt TEXT,
    importance_v1 REAL,            -- raw score from tools/osm_score.py
    importance_v1_uat_z REAL,      -- z-score within UAT
    UNIQUE (uat_siruta, name_normalized)
);
CREATE INDEX ix_osm_uat       ON osm_streets(uat_siruta);
CREATE INDEX ix_osm_namenorm  ON osm_streets(name_normalized);
CREATE INDEX ix_osm_corenorm  ON osm_streets(core_name_norm);

CREATE TABLE street_osm_matches (
    street_id INTEGER NOT NULL REFERENCES streets(id),
    osm_street_id INTEGER NOT NULL REFERENCES osm_streets(id),
    match_type TEXT NOT NULL,      -- 'exact_normalized' | 'fuzzy_core_name'
    confidence REAL NOT NULL,
    PRIMARY KEY (street_id, osm_street_id)
);
CREATE INDEX ix_match_osm     ON street_osm_matches(osm_street_id);

-- ===== Postal registry scaffolds (populated by tools/postal_*.py) =====
-- Source: data/reference/coduri-postale+/infocod-cu-siruta-mai-2016.xlsx
-- (Bucuresti + Localitati peste 50.000 loc sheets only — the sub-50.000 sheet
-- has no street-level columns, a structural gap this source can't close; see
-- CODE_SPEC §12). One row per (uat_siruta, name_normalized), same grain as
-- osm_streets/streets_dedup. `name` excludes the street-type prefix (the
-- source already separates Tip artera/Denumire artera), unlike osm_streets.
CREATE TABLE postal_streets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    source_sheet TEXT NOT NULL,           -- 'Bucuresti' | 'Localitati peste 50.000 loc'
    source_row_ids TEXT NOT NULL,         -- JSON array of contributing raw row indices
    judet_raw TEXT,
    localitate_raw TEXT,
    uat_siruta INTEGER,                   -- resolved & validated against streets.siruta; NULL = unresolved
    resolution_method TEXT NOT NULL,      -- 'direct_sector' | 'direct_sirsup' | 'name_match_exact'
                                           -- | 'name_match_fuzzy' | 'unresolved'
    resolution_confidence REAL NOT NULL,
    tip_artera_raw TEXT,
    street_type TEXT,
    name TEXT NOT NULL,
    name_normalized TEXT NOT NULL,
    title TEXT,
    rank TEXT,
    is_saint INTEGER NOT NULL DEFAULT 0,
    is_date INTEGER NOT NULL DEFAULT 0,
    is_numeric INTEGER NOT NULL DEFAULT 0,
    core_name TEXT,
    core_name_norm TEXT,
    core_name_norm_swapped TEXT,          -- 2-token reorder of core_name_norm, else NULL
    UNIQUE (uat_siruta, name_normalized)
);
CREATE INDEX ix_postal_uat      ON postal_streets(uat_siruta);
CREATE INDEX ix_postal_namenorm ON postal_streets(name_normalized);
CREATE INDEX ix_postal_corenorm ON postal_streets(core_name_norm);
CREATE INDEX ix_postal_swapped  ON postal_streets(core_name_norm_swapped);

CREATE TABLE street_postal_matches (
    street_id INTEGER NOT NULL REFERENCES streets(id),
    postal_street_id INTEGER NOT NULL REFERENCES postal_streets(id),
    match_type TEXT NOT NULL,      -- 'exact_normalized' | 'fuzzy_core_name' | 'reordered_core_name'
    confidence REAL NOT NULL,
    PRIMARY KEY (street_id, postal_street_id)
);
CREATE INDEX ix_pmatch_postal ON street_postal_matches(postal_street_id);

CREATE INDEX ix_streets_judet     ON streets(judet);
CREATE INDEX ix_streets_uat       ON streets(uat);
CREATE INDEX ix_streets_namenorm  ON streets(name_normalized);
CREATE INDEX ix_streets_corenorm  ON streets(core_name_norm);
CREATE INDEX ix_streets_type      ON streets(street_type);
CREATE INDEX ix_streets_flags     ON streets(is_saint, is_date, is_numeric);
-- Composite index allows the optimizer to push WHERE siruta=? into streets_dedup.
CREATE INDEX ix_streets_siruta_name ON streets(siruta, name_normalized);
-- Supports the cross-source anti-join in the external_corroboration_gap query
-- (docs/queries.sql) — without it, that query falls back to a slow scan.
CREATE INDEX ix_streets_siruta_corenorm ON streets(siruta, core_name_norm);
CREATE INDEX ix_aliases_norm      ON street_aliases(alias_normalized);

-- GROUP BY siruta (not uat) is intentional: 48 UAT names are shared across
-- multiple județe and would be incorrectly merged if grouped by name alone.
CREATE VIEW streets_dedup AS
SELECT MIN(id) AS id, judet, uat, siruta, street_type, name, name_normalized,
       title, rank, is_saint, is_date, is_numeric, core_name, core_name_norm
FROM streets
WHERE name_normalized != ''
GROUP BY siruta, name_normalized;

CREATE VIEW streets_classified_pct AS
-- One row per distinct core_name_norm (non-null).
-- classification: first matching bucket; NULL means unclassified.
-- street_count: number of (uat, name) pairs sharing this core_name_norm.
-- Query this view grouped by classification for curation coverage reports.
WITH base AS (
  SELECT core_name_norm,
    MAX(is_numeric) AS is_numeric,
    MAX(is_date)    AS is_date,
    MAX(is_saint)   AS is_saint,
    COUNT(*)        AS street_count
  FROM streets_dedup
  WHERE core_name_norm IS NOT NULL
  GROUP BY core_name_norm
)
SELECT
  b.core_name_norm,
  b.street_count,
  CASE
    WHEN b.is_numeric  = 1             THEN 'numeric'
    WHEN b.is_date     = 1             THEN 'date'
    WHEN b.is_saint    = 1             THEN 'saint'
    WHEN p.core_name_norm  IS NOT NULL THEN 'person'
    WHEN n.core_name_norm  IS NOT NULL THEN 'nature'
    WHEN c.core_name_norm  IS NOT NULL THEN 'category'
    WHEN pr.core_name_norm IS NOT NULL THEN 'place'
    ELSE NULL
  END AS classification
FROM base b
LEFT JOIN persons         p  ON p.core_name_norm  = b.core_name_norm
LEFT JOIN nature_terms    n  ON n.core_name_norm  = b.core_name_norm
LEFT JOIN name_categories c  ON c.core_name_norm  = b.core_name_norm
LEFT JOIN place_refs      pr ON pr.core_name_norm = b.core_name_norm;

-- Additive consolidation of all 3 street-name sources: registry (authoritative,
-- full attributes) + streets OSM/postal found that have no registry match.
-- Does not touch streets_dedup. Cross-source comparison must join on
-- core_name_norm, never name_normalized — osm_streets.name_normalized includes
-- the street-type prefix, postal's and the registry's don't. See CODE_SPEC §12.
CREATE VIEW streets_all_sources AS
SELECT 'registry' AS source, sd.id, sd.judet, sd.uat, sd.siruta, sd.street_type, sd.name,
       sd.name_normalized, sd.title, sd.rank, sd.is_saint, sd.is_date, sd.is_numeric,
       sd.core_name, sd.core_name_norm, NULL AS source_note
FROM streets_dedup sd

UNION ALL

SELECT 'osm', NULL, su.judet, su.uat, o.uat_siruta, o.street_type, o.name, o.name_normalized,
       o.title, o.rank, o.is_saint, o.is_date, o.is_numeric, o.core_name, o.core_name_norm,
       'highway=' || o.highway_class
FROM osm_streets o
LEFT JOIN (SELECT DISTINCT siruta, judet, uat FROM streets) su ON su.siruta = o.uat_siruta
WHERE NOT EXISTS (SELECT 1 FROM street_osm_matches m WHERE m.osm_street_id = o.id)

UNION ALL

SELECT 'postal', NULL, su.judet, su.uat, p.uat_siruta, p.street_type, p.name, p.name_normalized,
       p.title, p.rank, p.is_saint, p.is_date, p.is_numeric, p.core_name, p.core_name_norm,
       'sheet=' || p.source_sheet
FROM postal_streets p
LEFT JOIN (SELECT DISTINCT siruta, judet, uat FROM streets) su ON su.siruta = p.uat_siruta
WHERE p.uat_siruta IS NOT NULL
  AND NOT EXISTS (SELECT 1 FROM street_postal_matches m WHERE m.postal_street_id = p.id);
""")

# ---------- ingest ----------
wb = openpyxl.load_workbook(SRC, read_only=True)
ws = wb[wb.sheetnames[0]]
inserted = 0
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if row[0] is None: break
    if LIMIT and inserted >= LIMIT: break
    judet, uat, siruta = row[0], row[1], row[2]
    artera = row[10]
    if not artera: continue
    st, name, aliases = parse_artery(artera)
    feats = extract_features(name)
    cur = con.execute("""
        INSERT INTO streets (judet, uat, siruta, artera_raw, street_type, name,
            name_normalized, title, rank, is_saint, is_date, is_numeric,
            core_name, core_name_norm)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (judet, uat, siruta, artera, st, name,
          normalize_match(name), feats["title"], feats["rank"],
          feats["is_saint"], feats["is_date"], feats["is_numeric"],
          feats["core_name"], normalize_match(feats["core_name"]) if feats["core_name"] else None))
    sid = cur.lastrowid
    for a in aliases:
        con.execute("INSERT INTO street_aliases(street_id,alias,alias_normalized) VALUES (?,?,?)",
                    (sid, a, normalize_match(a)))
    inserted += 1

con.commit()
def n(sql): return con.execute(sql).fetchone()[0]
print(f"Inserted     {inserted}")
print(f"Aliases      {n('SELECT COUNT(*) FROM street_aliases')}")
print(f"Deduped      {n('SELECT COUNT(*) FROM streets_dedup')}")
print(f"Saints       {n('SELECT COUNT(*) FROM streets_dedup WHERE is_saint=1')}")
print(f"Dates        {n('SELECT COUNT(*) FROM streets_dedup WHERE is_date=1')}")
print(f"Numeric      {n('SELECT COUNT(*) FROM streets_dedup WHERE is_numeric=1')}")
print(f"With title   {n('SELECT COUNT(*) FROM streets_dedup WHERE title IS NOT NULL')}")
print(f"With rank    {n('SELECT COUNT(*) FROM streets_dedup WHERE rank IS NOT NULL')}")
con.close()
